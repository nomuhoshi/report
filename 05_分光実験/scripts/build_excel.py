# -*- coding: utf-8 -*-
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import ScatterChart, Reference, Series
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.worksheet.properties import PageSetupProperties

SRC = '/mnt/user-data/uploads/实验报告/bunkou.xlsx'
OUT = '/tmp/claude-0/-home-claude/c15e1565-df52-512f-94d3-ab60fda4dd6a/scratchpad/w2/bunkou.xlsx'

FN = 'Yu Gothic'
f_norm  = Font(name=FN, size=10)
f_bold  = Font(name=FN, size=10, bold=True)
f_in    = Font(name=FN, size=10, color='0000FF')     # 手稿から転記した入力値
f_title = Font(name=FN, size=14, bold=True)
f_note  = Font(name=FN, size=9, color='595959')
fill_head = PatternFill('solid', fgColor='D9E1F2')
fill_chk  = PatternFill('solid', fgColor='FFF2CC')   # 要確認
fill_pk   = PatternFill('solid', fgColor='E2EFDA')   # ピーク
thin = Side(style='thin', color='808080')
box  = Border(left=thin, right=thin, top=thin, bottom=thin)
ctr  = Alignment(horizontal='center', vertical='center', wrap_text=True)
left = Alignment(horizontal='left', vertical='center', wrap_text=True)
nowrap = Alignment(horizontal='left', vertical='center', wrap_text=False)

def put(ws, ref, val, font=f_norm, fill=None, fmt=None, al=ctr, border=True):
    c = ws[ref]; c.value = val; c.font = font; c.alignment = al
    if fill: c.fill = fill
    if fmt: c.number_format = fmt
    if border: c.border = box
    return c

def head(ws, row, cols, labels, h=30):
    for c, t in zip(cols, labels):
        put(ws, '%s%d' % (c, row), t, f_bold, fill_head)
    ws.row_dimensions[row].height = h

black = lambda: LineProperties(solidFill='000000', w=9525, prstDash='solid')
gray  = lambda: LineProperties(solidFill='808080', w=9525, prstDash='solid')

def make_chart(ws, series_defs, xspec, yspec, anchor, w=14, h=8.5):
    """series_defs: list of (xref, yref, title, marker, linecolor)"""
    ch = ScatterChart(); ch.style = 2; ch.title = None; ch.roundedCorners = False
    ch.height = h; ch.width = w
    ch.legend = None
    for xref, yref, title, marker, linecolor in series_defs:
        s = Series(yref, xref, title=title)
        if marker:
            s.marker.symbol = 'circle'; s.marker.size = 5
            s.marker.graphicalProperties = GraphicalProperties(solidFill=marker)
            s.marker.graphicalProperties.line.solidFill = marker
        else:
            s.marker.symbol = 'none'
        s.graphicalProperties = GraphicalProperties()
        if linecolor:
            s.graphicalProperties.line = LineProperties(solidFill=linecolor, w=9525, prstDash='solid')
        else:
            s.graphicalProperties.line.noFill = True
        ch.series.append(s)
    if len([1 for d in series_defs if d[2]]) > 1:
        ch.legend = ch.legend or None
    for ax, (mn, mx, mj, fmt, title) in ((ch.x_axis, xspec), (ch.y_axis, yspec)):
        ax.scaling.min = mn; ax.scaling.max = mx; ax.majorUnit = mj
        ax.majorTickMark = 'in'; ax.minorTickMark = 'none'
        ax.majorGridlines = None; ax.minorGridlines = None
        ax.crosses = 'min'; ax.delete = False
        ax.number_format = fmt
        if ax.numFmt is not None: ax.numFmt.sourceLinked = False
        ax.spPr = GraphicalProperties(ln=black())
        ax.title = title
    ch.plot_area.spPr = GraphicalProperties(noFill=True, ln=black())
    ch.graphical_properties = GraphicalProperties(noFill=True, ln=LineProperties(noFill=True))
    ws.add_chart(ch, anchor)
    return ch

# ====================================================== データ
# Cd 連続掃引（入射スリット 0.1 mm）。I は [V]。
CD_PASS1 = [(250,.0048),(255,.0050),(260,.0051),(265,.0051),(270,.0053),(275,.0059),(280,.0065),
            (282,.0063),(284,.0059),(286,.0058),(288,.0056),(290,.0055),(292,.0054),(294,.0052),(296,.0050)]
CD_PASS2 = [(220,.0051),(225,.0053),(230,.0053),(235,.0054),(240,.0055),(245,.0056),(250,.0057),
  (255,.0057),(260,.0058),(262,.0057),(264,.0057),(266,.0057),(268,.0058),(270,.0059),(272,.0058),
  (274,.0062),(275,.0064),(276,.0068),(277,.0069),(278,.0070),(279,.0070),(280,.0069),(281,.0067),
  (282,.0065),(283,.0063),(284,.0062),(285,.0061),(286,.0060),(287,.0059),(288,.0059),(289,.0058),
  (290,.0058),(292,.0057),(294,.0054),(296,.0051),(298,.0050),(300,.0049),(325,.0048),(350,.0049),
  (375,.0049),(400,.0049),(425,.0049),(450,.0049),(475,.0049),(500,.0047),(525,.0048),(550,.0048),
  (575,.0048),(600,.0050),(625,.0059),(626,.0062),(627,.0065),(628,.0069),(629,.0074),(630,.0080),
  (631,.0081),(632,.0083),(633,.0082),(634,.0089),(635,.0077),(636,.0074),(637,.0073),(638,.0070),
  (639,.0069),(640,.0069),(641,.0068),(642,.0067),(643,.0067),(644,.0066),(645,.0064),(646,.0064),
  (647,.0063),(648,.0062),(649,.0060),(650,.0060),(651,.0058),(652,.0056),(653,.0055),(654,.0053),
  (655,.0051),(656,.0050),(658,.0050),(660,.0050),(662,.0050),(664,.0050),(666,.0050),(668,.0051),
  (670,.0051),(672,.0051),(674,.0051),(676,.0051),(678,.0052),(680,.0052),(682,.0052),(684,.0052),
  (686,.0054),(688,.0053),(690,.0054),(695,.0056),(700,.0061),(701,.0064),(702,.0067),(703,.0069),
  (704,.0072),(705,.0075),(706,.0072),(707,.0071),(708,.0070),(710,.0067),(712,.0066),(714,.0065),
  (716,.0064),(718,.0063),(720,.0062),(722,.0061),(724,.0059),(726,.0056),(728,.0055),(730,.0055),
  (732,.0055),(734,.0057),(736,.0058),(738,.0059),(740,.0056),(741,.0054),(742,.0054),(743,.0054)]
# 3-3 入射スリット 1 mm（x=600 に戻して再掃引）
CD_SLIT1 = [(605,.0066),(610,.0072),(615,.0082),(620,.0123),(625,.0148),(630,.0159),(635,.0162),
            (640,.0156),(645,.0132),(650,.0104),(655,.0082),(660,.0068)]
# 3-4 白色 LED
LED = [(200,.0690),(225,.0815),(250,.0972),(275,.1134),(300,.1317),(325,.1485),(350,.1655),(375,.1795),
  (400,.1903),(425,.1966),(450,.1988),(475,.1965),(500,.1894),(525,.1784),(550,.1623),(575,.1412),
  (600,.1109),(625,.0778),(650,.0472),(675,.0312),(700,.0312),(702,.0316),(704,.0325),(706,.0330),
  (708,.0337),(710,.0344),(712,.0352),(714,.0363),(716,.0375),(718,.0388),(720,.0402),(722,.0420),
  (724,.0438),(726,.0456),(728,.0475),(730,.0497),(732,.0522),(734,.0546),(736,.0570),(738,.0593),
  (740,.0620),(742,.0645),(744,.0673),(746,.0700),(748,.0730),(750,.0761),(752,.0796),(754,.0837),
  (756,.0875),(758,.0922),(760,.0969),(762,.1018),(764,.1072),(766,.1132),(768,.1188),(770,.1233),
  (772,.1278),(774,.1313),(776,.1338),(778,.1352),(780,.1349),(782,.1328),(785,.1274),(790,.1123),
  (795,.0941),(800,.0753),(805,.0594),(810,.0466),(815,.0367),(820,.0296),(825,.0236),(830,.0193),
  (835,.0156),(840,.0132)]

wb = openpyxl.load_workbook(SRC)      # week_1 はそのまま残す

# =============================================== 実験1_波長
ws1 = wb.create_sheet('実験1_波長')
ws1.sheet_view.showGridLines = False
put(ws1,'B1','実験1　各種ランプの輝線の波長（反射型回折格子・1次回折光）',f_title,al=nowrap,border=False)
put(ws1,'B2','式: tanθ = (x0 + x1) / L ，　d sinθ = mλ（m = 1）。青字＝手稿から転記した入力値，黒字＝Excel公式，黄色＝要確認。',f_note,al=nowrap,border=False)
head(ws1,4,'BCD',['定数','値','単位'],18)
ws1.merge_cells('E4:H4'); put(ws1,'E4','出典',f_bold,fill_head)
for col in 'FGH': ws1['%s4'%col].border=box
consts=[('格子定数 d','=10^6/1200','nm','講義 [装置/器具] 2-2: 刻線数 1200 lines/mm → d = 1/1200 mm',False),
        ('回折格子–スクリーン間距離 L',150,'mm','手稿 第1週 p.2（L → 150 mm）',True),
        ('x0（光軸の位置）',50,'mm','手稿 第1週 p.2（x0 → 50 mm）',True),
        ('回折次数 m',1,'－','手稿 第1週 p.1（m = +1 のみ）',True)]
for i,(n,v,u,s,inp) in enumerate(consts):
    r=5+i
    put(ws1,'B%d'%r,n,al=left); put(ws1,'C%d'%r,v,f_in if inp else f_norm,fmt='0.0' if i<3 else '0')
    put(ws1,'D%d'%r,u); ws1.merge_cells('E%d:H%d'%(r,r)); put(ws1,'E%d'%r,s,f_note,al=left)
    for col in 'FGH': ws1['%s%d'%(col,r)].border=box
HR=11
head(ws1,HR,'BCDEFGHIJKLMNO',['ランプ','輝線の色','x0 [mm]','L [mm]','x1 左 [mm]','x1 中 [mm]','x1 右 [mm]',
    'x = x0+x1 [mm]','θ [deg]','λ [nm]','文献値 [nm]','λ − 文献値 [nm]','相対誤差 [%]','出典・備考'],42)
rows=[
 ('Na','オレンジ（報告書に使用）','=$C$7','=$C$6',95,97,99,'=AVERAGE(588.995,589.592)',
  '手稿 第1週 p.2。報告書にはこの組（95 / 97 / 99）を使う。文献値は D線2本(588.995, 589.592 nm)の平均',{}),
 ('Na','オレンジ（参考・不使用）','=$C$7','=$C$6',95.2,98.1,100.4,'=AVERAGE(588.995,589.592)',
  '手稿 第1週 p.3 の再測定値。参考として残すが報告書には使わない',{}),
 ('Hg','オレンジ','=$C$7','=$C$6',87.5,89.0,90.9,579.065,'手稿 第1週 p.4。文献値 579.065 nm（講義 付録）',{'L':1}),
 ('Hg','緑','=$C$7','=$C$6',74.2,75.5,77.0,546.074,'手稿 第1週 p.4。文献値 546.074 nm（講義 付録）',{}),
 ('Hg','青','=$C$7','=$C$6',37.1,38.2,39.2,435.835,'手稿 第1週 p.4。文献値 435.835 nm（講義 付録）',{}),
 ('Cd','赤','=$C$7','=$C$6',None,12.5,None,643.847,
  '手稿 第1週 p.5「真ん中 12.5 mm」。左右は読み取りにくく中心のみ。x0・L は Na と同じとした。計算値は文献値の約半分で外れ値',{}),
 ('Cd','緑','=$C$7','=$C$6',52.0,54.5,57.0,508.582,'手稿 第1週 p.5。x0・L は Na と同じとした',{}),
 ('Cd','水色','=$C$7','=$C$6',35.5,37.0,40.5,479.992,'手稿 第1週 p.5。x0・L は Na と同じとした',{}),
 ('Cd','青','=$C$7','=$C$6',35.6,None,None,467.815,
  '手稿 第1週 p.5「青 左 35.6」のみ。光が細く中・右は読み取れなかった。λ は左端の値で計算。第2週の結果より，この輝線は 466.235 と 467.815 nm の未分解の重なりと分かった',{'L':1}),
]
for i,(lamp,col,x0,L,xl,xm,xr,lit,note,fl) in enumerate(rows):
    r=HR+1+i
    put(ws1,'B%d'%r,lamp); put(ws1,'C%d'%r,col)
    put(ws1,'D%d'%r,x0,fmt='0.0'); put(ws1,'E%d'%r,L,fmt='0.0')
    for c,v in zip('FGH',(xl,xm,xr)): put(ws1,'%s%d'%(c,r),v,f_in,fmt='0.0')
    pc = 'F' if (lamp=='Cd' and col=='青') else 'G'
    put(ws1,'I%d'%r,'=IF(COUNT(D%d,E%d,%s%d)=3,D%d+%s%d,"")'%(r,r,pc,r,r,pc,r),fmt='0.0')
    put(ws1,'J%d'%r,'=IF(I%d="","",DEGREES(ATAN(I%d/E%d)))'%(r,r,r),fmt='0.00')
    put(ws1,'K%d'%r,'=IF(J%d="","",$C$5*SIN(RADIANS(J%d))/$C$8)'%(r,r),fmt='0.0')
    put(ws1,'L%d'%r,lit,f_in if not isinstance(lit,str) else f_norm,fmt='0.000')
    put(ws1,'M%d'%r,'=IF(K%d="","",K%d-L%d)'%(r,r,r),fmt='+0.0;−0.0;0.0')
    put(ws1,'N%d'%r,'=IF(K%d="","",ABS(M%d)/L%d*100)'%(r,r,r),fmt='0.0')
    put(ws1,'O%d'%r,note,f_note,al=left)
    for c in ('D','E','G','L'):
        if fl.get(c): ws1['%s%d'%(c,r)].fill=fill_chk
    ws1.row_dimensions[r].height=40
for col,w in zip('ABCDEFGHIJKLMNO',[2,26,20,9,9,11,11,11,14,10,10,12,13,11,62]):
    ws1.column_dimensions[col].width=w
ws1.freeze_panes='D%d'%(HR+1)

# =============================================== 実験2_屈折率
ws2 = wb.create_sheet('実験2_屈折率')
ws2.sheet_view.showGridLines = False
put(ws2,'B1','実験2　最小偏角 δ0 とプリズム（ガラス）の屈折率',f_title,al=nowrap,border=False)
put(ws2,'B2','式: tanδ0 = L / x ，　n = sin{(δ0 + α)/2} / sin(α/2)（講義 式(4)）。青字＝手稿から転記，黒字＝公式。',f_note,al=nowrap,border=False)
put(ws2,'B4','頂角 α',al=left); put(ws2,'C4',60,f_in,fmt='0'); put(ws2,'D4','deg')
ws2.merge_cells('E4:J4'); put(ws2,'E4','講義 [装置/器具] 2-1: プリズムは正三角形，頂角 α = 60°',f_note,al=left)
for c in 'FGHIJ': ws2['%s4'%c].border=box
put(ws2,'B6','色',f_bold,fill_head); put(ws2,'C6','波長 λ\n[nm]',f_bold,fill_head)
ws2.merge_cells('B6:B7'); ws2.merge_cells('C6:C7'); ws2['B7'].border=box; ws2['C7'].border=box
groups=[('D','E','F',150),('G','H','I',175),('J','K','L',200)]
for a,b,c,Lv in groups:
    ws2.merge_cells('%s6:%s6'%(a,c)); put(ws2,'%s6'%a,Lv,f_in,fill_head,fmt='"L = "0" mm"')
    ws2['%s6'%b].border=box; ws2['%s6'%c].border=box
    put(ws2,'%s7'%a,'x [mm]',f_bold,fill_head); put(ws2,'%s7'%b,'δ0 [deg]',f_bold,fill_head); put(ws2,'%s7'%c,'n',f_bold,fill_head)
ws2.merge_cells('M6:M7'); put(ws2,'M6','n 平均\n（3回）',f_bold,fill_head); ws2['M7'].border=box
ws2.merge_cells('N6:N7'); put(ws2,'N6','出典・備考',f_bold,fill_head); ws2['N7'].border=box
xs={'D':[136,128,127,118],'G':[155,151,149,143],'J':[188,181,176,170]}
colors=[('赤',643.847,'Cd 赤 643.847 nm（講義 付録）。x は手稿 第1週 p.6'),
        ('緑',508.582,'Cd 緑 508.582 nm。x は手稿 第1週 p.6'),
        ('水色',479.992,'Cd 水色 479.992 nm。x は手稿 第1週 p.6'),
        ('青',467.815,'Cd 青 467.815 nm（466.235 nm との未分解の重なり）。x は手稿 第1週 p.6')]
for i,(nm,lf,note) in enumerate(colors):
    r=8+i
    put(ws2,'B%d'%r,nm); put(ws2,'C%d'%r,lf,f_in,fmt='0.0')
    for a,b,c,Lv in groups:
        put(ws2,'%s%d'%(a,r),xs[a][i],f_in,fmt='0')
        put(ws2,'%s%d'%(b,r),'=DEGREES(ATAN(%s$6/%s%d))'%(a,a,r),fmt='0.00')
        put(ws2,'%s%d'%(c,r),'=SIN(RADIANS((%s%d+$C$4)/2))/SIN(RADIANS($C$4/2))'%(b,r),fmt='0.0000')
    put(ws2,'M%d'%r,'=AVERAGE(F%d,I%d,L%d)'%(r,r,r),fmt='0.0000')
    put(ws2,'N%d'%r,note,f_note,al=left)
    ws2.row_dimensions[r].height=24
make_chart(ws2,[(Reference(ws2,min_col=3,min_row=8,max_row=11),
                 Reference(ws2,min_col=13,min_row=8,max_row=11),None,'000000',None)],
           (440,680,40,'0','波長 λ [nm]'),(1.60,1.68,0.02,'0.00','屈折率 n'),'B14')
put(ws2,'B32','※ 図名は図の下に付ける。横軸 440–680 nm，縦軸 1.60–1.68 に固定。',f_note,al=nowrap,border=False)
for col,w in zip('ABCDEFGHIJKLMN',[2,10,10,8,9,9,8,9,9,8,9,9,11,52]):
    ws2.column_dimensions[col].width=w
ws2.row_dimensions[6].height=28; ws2.row_dimensions[7].height=20

# =============================================== 実験3-2 波長校正
w32 = wb.create_sheet('実験3-2_波長校正')
w32.sheet_view.showGridLines = False
put(w32,'B1','実験3-2　Cd ランプのスペクトル測定と波長校正',f_title,al=nowrap,border=False)
put(w32,'B2','x はマイクロメータ読み（1 目盛 = 0.01 mm，1 回転 = 0.5 mm）。回転ステージ角は θr = k x + α と書けるので，式(10) に代入すると λ = A x² + B x + C（講義 式(10)＋授業の板書）。青字＝手稿から転記，黒字＝Excel公式，緑＝ピーク，黄色＝要確認。',f_note,al=nowrap,border=False)
put(w32,'B3','入射スリット幅 0.1 mm（手稿 第2週 p.4 見出し「幅 0.1mm」）。λ が x とともに減少することに注意。',f_note,al=nowrap,border=False)

head(w32,5,'BCDEFG',['輝線の色','x ピーク','I ピーク [V]','文献波長 [nm]','校正に使用','備考'],20)
pk_rows=[('赤','=AVERAGE(278,279)',.0070,643.847,'○','x=278,279 が同値（0.0070）のため中点を採用'),
         ('緑',634,.0089,508.582,'○','スリットで緑色を目視確認（手稿の「←緑」の書き込み）。4本中で最も強い → 実験3-3 に使用'),
         ('水色',705,.0075,479.992,'○',''),
         ('青',738,.0059,467.815,'×','466.235 nm と 467.815 nm（間隔 1.58 nm）の未分解の重なり。分解能 4.4 nm より狭いため分離できない。校正には使わず検証のみ')]
for i,(c,xp,ip,lit,use,note) in enumerate(pk_rows):
    r=6+i
    put(w32,'B%d'%r,c); put(w32,'C%d'%r,xp,f_in,fmt='0.0',fill=fill_pk)
    put(w32,'D%d'%r,ip,f_in,fmt='0.0000',fill=fill_pk)
    put(w32,'E%d'%r,lit,f_in,fmt='0.000'); put(w32,'F%d'%r,use)
    put(w32,'G%d'%r,note,f_note,al=left); w32.row_dimensions[r].height=34

put(w32,'B11','校正曲線　λ = A x² + B x + C',f_bold,al=nowrap,border=False)
for lab,cell,idx,fmt in (('A [nm/目盛²]','C12',1,'0.000E+00'),('B [nm/目盛]','C13',2,'0.00000'),('C [nm]','C14',3,'0.000')):
    r=int(cell[1:]); put(w32,'B%d'%r,lab,al=left)
    put(w32,cell,'=INDEX(LINEST($E$6:$E$8,$C$6:$C$8^{1,2}),1,%d)'%idx,fmt=fmt)
ws_note=['LINEST による最小二乗フィット（赤・緑・水色の3点，未知数3個なので厳密に一致）。',
         '','']
for i,t in enumerate(ws_note):
    if t: put(w32,'D%d'%(12+i),t,f_note,al=nowrap,border=False)
put(w32,'B15','検証',f_bold,al=left)
checks=[('x = 738 の λ','=$C$12*738^2+$C$13*738+$C$14','nm','実測の弱いピーク。467.815 と 466.235 nm の中間に出る → 未分解の重なりの証拠'),
        ('λ = 467.815 nm の x','=(-$C$13-SQRT($C$13^2-4*$C$12*($C$14-467.815)))/(2*$C$12)','目盛',''),
        ('λ = 466.235 nm の x','=(-$C$13-SQRT($C$13^2-4*$C$12*($C$14-466.235)))/(2*$C$12)','目盛','')]
for i,(lab,f,u,note) in enumerate(checks):
    r=16+i; put(w32,'B%d'%r,lab,al=left); put(w32,'C%d'%r,f,fmt='0.0'); put(w32,'D%d'%r,u)
    if note: put(w32,'E%d'%r,note,f_note,al=nowrap,border=False)

# data table
DR=21
head(w32,DR,'BCD',['x [目盛]','I [V]','λ [nm]'],18)
put(w32,'E%d'%DR,'掃引',f_bold,fill_head)
r=DR+1
for x,v in CD_PASS1:
    put(w32,'B%d'%r,x,f_in,fmt='0'); put(w32,'C%d'%r,v,f_in,fmt='0.0000')
    put(w32,'D%d'%r,'=$C$12*B%d^2+$C$13*B%d+$C$14'%(r,r),fmt='0.0')
    put(w32,'E%d'%r,'1回目（赤線のみ）',f_note); r+=1
P2=r
for x,v in CD_PASS2:
    put(w32,'B%d'%r,x,f_in,fmt='0'); put(w32,'C%d'%r,v,f_in,fmt='0.0000')
    put(w32,'D%d'%r,'=$C$12*B%d^2+$C$13*B%d+$C$14'%(r,r),fmt='0.0')
    put(w32,'E%d'%r,'2回目（本掃引）',f_note); r+=1
P2E=r-1
for col,w in zip('ABCDEFGHIJ',[2,16,12,12,13,12,62,3,12,12]): w32.column_dimensions[col].width=w

# 図: Cd スペクトル（横軸 波長，縦軸 出力電圧）
make_chart(w32,[(Reference(w32,min_col=4,min_row=P2,max_row=P2E),
                 Reference(w32,min_col=3,min_row=P2,max_row=P2E),None,None,'000000')],
           (440,680,40,'0','波長 λ [nm]'),(0.004,0.010,0.001,'0.000','出力電圧 I [V]'),'L21')
# 図: 波長校正曲線（横軸 ステージ位置，縦軸 波長）
FC=DR+1
put(w32,'I%d'%(DR-2),'校正曲線描画用',f_bold,al=nowrap,border=False)
head(w32,DR,'IJ',['x [目盛]','λ_fit [nm]'],18)
rr=FC
for xv in range(200,861,20):
    put(w32,'I%d'%rr,xv,fmt='0'); put(w32,'J%d'%rr,'=$C$12*I%d^2+$C$13*I%d+$C$14'%(rr,rr),fmt='0.0'); rr+=1
make_chart(w32,[(Reference(w32,min_col=9,min_row=FC,max_row=rr-1),
                 Reference(w32,min_col=10,min_row=FC,max_row=rr-1),'近似曲線',None,'000000'),
                (Reference(w32,min_col=3,min_row=6,max_row=8),
                 Reference(w32,min_col=5,min_row=6,max_row=8),'Cd 輝線','FF0000',None)],
           (200,860,100,'0','マイクロメータ読み x [目盛]'),(400,700,50,'0','波長 λ [nm]'),'L39')

w32.freeze_panes='B%d'%(DR+1)

# =============================================== 実験3-3 分解能
w33 = wb.create_sheet('実験3-3_分解能')
w33.sheet_view.showGridLines = False
put(w33,'B1','実験3-3　スペクトル分解能測定（Cd 緑線 508.582 nm）',f_title,al=nowrap,border=False)
put(w33,'B2','実験3-2 で最も強かった緑線について，入射スリット幅を変えて同じ輝線を掃引した。半値幅は背景を差し引いた最大値の 1/2 の高さでの幅（線形内挿）。λ は実験3-2 の校正曲線で換算。',f_note,al=nowrap,border=False)
put(w33,'B3','黄色＝要確認。1 mm のデータは x=605〜660 しか無く，背景（遠方のベースライン）が測られていない。',f_note,al=nowrap,border=False)

# scan A table
head(w33,5,'BC',['x [目盛]','I [V]（0.1 mm）'],18)
A0=6; r=A0
sub=[d for d in CD_PASS2 if 620<=d[0]<=655]
for x,v in sub:
    put(w33,'B%d'%r,x,f_in,fmt='0'); put(w33,'C%d'%r,v,f_in,fmt='0.0000'); r+=1
AE=r-1
head(w33,5,'EF',['x [目盛]','I [V]（1 mm）'],18)
B0=6; r=B0
for x,v in CD_SLIT1:
    put(w33,'E%d'%r,x,f_in,fmt='0'); put(w33,'F%d'%r,v,f_in,fmt='0.0000'); r+=1
BE=r-1

# analysis block
put(w33,'H5','解析',f_bold,al=left,border=False)
head(w33,6,'HIJ',['項目','スリット 0.1 mm','スリット 1 mm'],18)
def rowdef(r,lab,fa,fb,fmt,note=None,chk=False):
    put(w33,'H%d'%r,lab,al=left); put(w33,'I%d'%r,fa,fmt=fmt); put(w33,'J%d'%r,fb,fmt=fmt)
    if chk: w33['I%d'%r].fill=fill_chk; w33['J%d'%r].fill=fill_chk
    if note: put(w33,'K%d'%r,note,f_note,al=nowrap,border=False)
rowdef(7,'背景 I_b [V]',0.0049,0.0067,'0.0000','0.1 mm は x=500〜600 の平坦部。1 mm は掃引両端(605, 660)の平均 → 遠方の背景は未測定なので暫定値',chk=True)
w33['I7'].font=f_in; w33['J7'].font=f_in
rowdef(8,'ピーク I_p [V]','=MAX($C$%d:$C$%d)'%(A0,AE),'=MAX($F$%d:$F$%d)'%(B0,BE),'0.0000')
rowdef(9,'ピーク位置 x_p [目盛]','=INDEX($B$%d:$B$%d,MATCH(I8,$C$%d:$C$%d,0))'%(A0,AE,A0,AE),
         '=INDEX($E$%d:$E$%d,MATCH(J8,$F$%d:$F$%d,0))'%(B0,BE,B0,BE),'0')
rowdef(10,'半値 I_h [V]','=I7+(I8-I7)/2','=J7+(J8-J7)/2','0.0000')
# interpolation brackets (explicit)
putn=lambda ref,v,note=None: (put(w33,ref,v,f_note,al=left,border=False))
rowdef(11,'左交点 x_L [目盛]',
       '=B%d+(I10-C%d)*(B%d-B%d)/(C%d-C%d)'%(A0+2,A0+2,A0+3,A0+2,A0+3,A0+2),
       '=E%d+(J10-F%d)*(E%d-E%d)/(F%d-F%d)'%(B0+2,B0+2,B0+3,B0+2,B0+3,B0+2),'0.0',
       '線形内挿の区間：0.1mm は x=627→628，1mm は x=615→620。交点は読み取りの刻み幅のため ±1 目盛（約 0.4 nm）の不確かさがある')
rowdef(12,'右交点 x_R [目盛]',
       '=B%d+(I10-C%d)*(B%d-B%d)/(C%d-C%d)'%(A0+13,A0+13,A0+14,A0+13,A0+14,A0+13),
       '=E%d+(J10-F%d)*(E%d-E%d)/(F%d-F%d)'%(B0+8,B0+8,B0+9,B0+8,B0+9,B0+8),'0.0',
       '線形内挿の区間：0.1mm は x=638→639，1mm は x=645→650。0.1mm では x=639 と 640 が同値のため右交点には ±1 目盛の幅がある')
rowdef(13,'半値幅 Δx [目盛]','=I12-I11','=J12-J11','0.0')
cal="'実験3-2_波長校正'!"
rowdef(14,'半値幅 Δλ [nm]',
       '=ABS((%sC$12*I11^2+%sC$13*I11+%sC$14)-(%sC$12*I12^2+%sC$13*I12+%sC$14))'%(cal,cal,cal,cal,cal,cal),
       '=ABS((%sC$12*J11^2+%sC$13*J11+%sC$14)-(%sC$12*J12^2+%sC$13*J12+%sC$14))'%(cal,cal,cal,cal,cal,cal),'0.00')
rowdef(15,'ピーク波長 [nm]',
       '=%sC$12*I9^2+%sC$13*I9+%sC$14'%(cal,cal,cal),
       '=%sC$12*J9^2+%sC$13*J9+%sC$14'%(cal,cal,cal),'0.0')

put(w33,'H17','まとめ（報告書の表）',f_bold,al=left,border=False)
head(w33,18,'HIJ',['入射スリット幅 [µm]','ピーク値 [V]','半値幅 [nm]'],18)
put(w33,'H19',100,f_in,fmt='0'); put(w33,'I19','=I8',fmt='0.0000'); put(w33,'J19','=I14',fmt='0.00')
put(w33,'H20',1000,f_in,fmt='0'); put(w33,'I20','=J8',fmt='0.0000'); put(w33,'J20','=J14',fmt='0.00')
put(w33,'K19','講義は 200〜1000 µm の範囲を指示しているが，実際は 100 µm と 1000 µm の 2 点。報告書にその旨を書く。',f_note,al=nowrap,border=False)

make_chart(w33,[(Reference(w33,min_col=8,min_row=19,max_row=20),
                 Reference(w33,min_col=10,min_row=19,max_row=20),None,'000000',None)],
           (0,1200,200,'0','入射スリット幅 [µm]'),(0,14,2,'0','半値幅 Δλ [nm]'),'H23')
put(w33,'B%d'%(AE+3),'※ 2 点しかないため直線の当てはめは行わない。',f_note,al=nowrap,border=False)
for col,w in zip('ABCDEFGHIJK',[2,11,14,3,11,14,3,22,16,16,70]): w33.column_dimensions[col].width=w

# =============================================== 実験3-4 白色LED
w34 = wb.create_sheet('実験3-4_白色LED')
w34.sheet_view.showGridLines = False
put(w34,'B1','実験3-4　白色 LED のスペクトル測定',f_title,al=nowrap,border=False)
put(w34,'B2','実験3-2 の校正曲線 λ = A x² + B x + C で x を波長に換算した。入射スリット幅 0.1 mm。',f_note,al=nowrap,border=False)
put(w34,'B3','講義は 350〜800 nm の測定を指示しているが，実測範囲は下の「測定範囲」の通り。報告書にその旨を書く。',f_note,al=nowrap,border=False)
DR4=9
head(w34,DR4,'BCD',['x [目盛]','I [V]','λ [nm]'],18)
r=DR4+1; L0=r
for x,v in LED:
    put(w34,'B%d'%r,x,f_in,fmt='0'); put(w34,'C%d'%r,v,f_in,fmt='0.0000')
    put(w34,'D%d'%r,'=%sC$12*B%d^2+%sC$13*B%d+%sC$14'%(cal,r,cal,r,cal),fmt='0.0'); r+=1
LE=r-1
put(w34,'F5','解析',f_bold,al=left,border=False)
an=[('測定範囲 [nm]','=TEXT(MIN(D%d:D%d),"0")&" 〜 "&TEXT(MAX(D%d:D%d),"0")'%(L0,LE,L0,LE),'@',''),
    ('発光帯ピーク I [V]','=MAX(C%d:C%d)'%(L0,LE),'0.0000','蛍光体による広帯域'),
    ('発光帯ピーク波長 [nm]','=INDEX(D%d:D%d,MATCH(MAX(C%d:C%d),C%d:C%d,0))'%(L0,LE,L0,LE,L0,LE),'0.0',''),
    ('青色ピーク I [V]','=MAX(C%d:C%d)'%(L0+40,LE),'0.0000','LED チップ自身の発光（x≧740 の範囲で探索）'),
    ('青色ピーク波長 [nm]','=INDEX(D%d:D%d,MATCH(MAX(C%d:C%d),C%d:C%d,0))'%(L0+40,LE,L0+40,LE,L0+40,LE),'0.0','')]
for i,(lab,f,fmt,note) in enumerate(an):
    r2=6+i; put(w34,'F%d'%r2,lab,al=left); put(w34,'G%d'%r2,f,fmt=fmt)
    if note: put(w34,'H%d'%r2,note,f_note,al=nowrap,border=False)
make_chart(w34,[(Reference(w34,min_col=4,min_row=L0,max_row=LE),
                 Reference(w34,min_col=3,min_row=L0,max_row=LE),None,None,'000000')],
           (420,680,40,'0','波長 λ [nm]'),(0,0.22,0.02,'0.00','出力電圧 I [V]'),'F13',w=16,h=9.5)
for col,w in zip('ABCDEFGH',[2,11,12,12,3,24,14,56]): w34.column_dimensions[col].width=w
w34.freeze_panes='B%d'%(DR4+1)

# =============================================== 説明
ws0 = wb.create_sheet('説明')
ws0.sheet_view.showGridLines = False
put(ws0,'B1','分光実験　データ整理（第1週：実験1・2／第2週：実験3-2・3-3・3-4）',f_title,al=nowrap,border=False)
lines=[('■ シート構成',1),
 ('・実験1_波長：Na・Hg・Cd の輝線から回折角 θ と波長 λ を計算し，文献値と比較',0),
 ('・実験2_屈折率：Cd 4色の最小偏角 δ0 → 屈折率 n，3回の平均，波長–屈折率のグラフ',0),
 ('・実験3-2_波長校正：Cd 掃引データ，4本のピーク同定，校正曲線 λ = A x² + B x + C，Cd スペクトルと校正曲線のグラフ',0),
 ('・実験3-3_分解能：緑線をスリット 0.1 mm / 1 mm で掃引，ピーク値と半値幅',0),
 ('・実験3-4_白色LED：LED 掃引データを波長に換算，スペクトルのグラフ',0),
 ('・week_1：元のデータ（変更していません）',0),
 ('■ 色のルール',1),
 ('・青字＝手稿から転記した測定値／定数　・黒字＝Excel の公式　・緑＝ピーク　・黄色背景＝要確認',0),
 ('■ 主な結果',1),
 ('・校正曲線は 赤 643.847 / 緑 508.582 / 水色 479.992 nm の3点から決定。白色 LED の2つのピークが 580 nm（蛍光体）と 450 nm（青色 LED チップ）になり，校正に使っていないデータで裏が取れている。',0),
 ('・λ は x の増加とともに減少する。',0),
 ('・分解能は 4.4 nm（スリット 0.1 mm）。Cd の 466.235 と 467.815 nm は間隔 1.58 nm なので分離できず，x≈738 の1つのピークとして現れる。',0),
 ('■ 講義の指示と実際の差（報告書に書くこと）',1),
 ('① 実験3-2 は「全ての輝線に対して3回行って平均」だが，実際は連続掃引（赤線のみ2回，他は1回）。データ点数が多く，所要時間の都合で繰り返せなかったと考えられる。',0),
 ('② 実験3-3 のスリット幅は「200〜1000 µm」の指示だが，実際は 100 µm と 1000 µm の2点。',0),
 ('③ 実験3-4 の測定範囲は「350〜800 nm」の指示だが，実測範囲は実験3-4 シートの「測定範囲」を参照。',0),
 ('④ 課題2 は Na・Cd・Hg の輝線幅と白色 LED を比べるが，実験3 で測ったのは Cd のみ。Na・Hg の輝線幅のデータはない。',0),
 ('■ 要確認（黄色のセル）',1),
 ('・実験3-3 の背景 I_b：1 mm の掃引は x=605〜660 しかなく遠方の背景が測られていない。暫定値を入れてあるので，半値幅 11.7 nm には不確かさがある。',0)]
r=3
for t,b in lines:
    ws0.merge_cells('B%d:L%d'%(r,r))
    put(ws0,'B%d'%r,t,f_bold if b else f_norm,al=left,border=False)
    ws0.row_dimensions[r].height=18 if (b or len(t)<50) else 32
    r+=1
ws0.column_dimensions['A'].width=2; ws0.column_dimensions['B'].width=16
for col in 'CDEFGHIJKL': ws0.column_dimensions[col].width=14

wb.move_sheet('説明', offset=-(len(wb.sheetnames)-1))
order=['説明','実験1_波長','実験2_屈折率','実験3-2_波長校正','実験3-3_分解能','実験3-4_白色LED','week_1']
wb._sheets=[wb[n] for n in order]
for w in wb:
    w.page_setup.orientation='landscape'; w.page_setup.fitToWidth=1; w.page_setup.fitToHeight=0
    w.sheet_properties.pageSetUpPr=PageSetupProperties(fitToPage=True)
wb.active=0
wb.save(OUT); print('saved', OUT)
